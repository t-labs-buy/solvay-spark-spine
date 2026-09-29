"""The BPML process house export -> Markdown, one section per process.

`BPML_Process.xlsx` is the Signavio export of Solvay's process house: one row
per process, 17 columns. Converted as a table (xlsx_tables.py) it made a poor
corpus document. One activity list, written with bare `|` separators, split
into 119 cells, so the whole table was padded to 119 columns and every chunk
carried the header, ~100 empty cells and a 119-cell separator row -- around
three quarters of each chunk was pipes. The embedding model saw the same block
in 1,382 chunks, and only 397 rows kept their values under the right header.

So this module writes it as records instead. Every process is a heading
("4.5.1 Create Sales Order") nested under its parent, which puts the hierarchy
in each chunk's heading path; its fields follow as a short list, then the
description and the activities. Metadata nobody asks about (author, release
and validity dates) is left out, and so are rows with no content of their own.

The format is also read back. `parse()` is how bpml.py builds the process
hierarchy InsightLens and the Fit-Gap Copilot scope with, from this document's
chunks in the corpus; `hierarchy()` is how the knowledge graph chains a cited
process, numbered or lettered, up to its value chain.

    .venv/bin/python -m backend.ingestion.bpml_markdown            # rebuild knowledge_base/BPML_Process_xlsx.md
    .venv/bin/python -m backend.ingestion.bpml_markdown --index    # ... and re-index it
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path

from backend.core.paths import ROOT

SOURCE = ROOT / "knowledge_base" / "BPML_Process.xlsx"
TARGET = ROOT / "knowledge_base" / "BPML_Process_xlsx.md"
TITLE = "BPML Process House"

# "4.5.1.4  Create Sales Order" -> code, name. A level-1 process is written
# "4.0"; every deeper level drops the trailing zero and adds a segment.
CODE_LINE = re.compile(r"^(\d+(?:\.\d+)+)\s+(.*)$")

COLUMNS = {
    "name": "Process Name", "path": "Root Path", "description": "Description",
    "accountable": "Accountable Organization", "process_type": "Process Type",
    "status": "Status", "inputs": "Inputs", "outputs": "Outputs",
    "roles": "Roles (Responsible)", "activities": "Activities",
}
# The field list under each heading: (label, attribute). parse() reads the
# same labels back, so a label changed here is a label changed there.
FIELDS = [
    ("Process type", "process_type"), ("Status", "status"),
    ("Accountable organisation", "accountable"), ("Roles", "roles"),
    ("Inputs", "inputs"), ("Outputs", "outputs"),
]
# Root Path segments every process shares; they say nothing about any one.
_PATH_PREFIX = {"Process House", "EtE Processes"}
_OFFSET = re.compile(r":?\s*«position-offset:(\d+)»")
_PLACEHOLDERS = {"", "To be Deleted", "Sub process"}


def level_of(code: str) -> int:
    """4.0 -> 1, 4.5 -> 2, 4.5.1 -> 3, 4.5.1.4 -> 4."""
    parts = code.split(".")
    if len(parts) == 2 and parts[1] == "0":
        return 1
    return len(parts)


def parent_of(code: str) -> str | None:
    parts = code.split(".")
    if level_of(code) == 1:
        return None
    if len(parts) == 2:           # 4.5 -> 4.0
        return f"{parts[0]}.0"
    return ".".join(parts[:-1])   # 4.5.1.4 -> 4.5.1


def sort_key(code: str) -> tuple:
    """Numeric sort, so 4.10 follows 4.9 instead of 4.1."""
    return tuple(int(x) if x.isdigit() else 0 for x in code.split("."))


def _fix(value: object) -> str:
    """Undo the export's mojibake. Its text is Windows-1252 read as Mac Roman:
    "SÈbastien", "companyís", "ìNew Initiationî", a no-break space as "†".
    Every non-ASCII character in the export is garbled this way, so the whole
    text is converted back; a string that does not round-trip is kept."""
    if value is None:
        return ""
    text = str(value)
    try:
        text = text.encode("mac_roman").decode("cp1252")
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return text.replace(" ", " ").strip()


@dataclass
class Record:
    code: str  # the BPML code, or "" for a BPMN object row ("PE-040-030 ...")
    name: str
    path: list[str] = field(default_factory=list)  # Root Path, below the house
    description: str = ""
    process_type: str = ""
    status: str = ""
    accountable: str = ""
    roles: str = ""
    inputs: str = ""
    outputs: str = ""
    activities: list[str] = field(default_factory=list)
    synthesised: bool = False  # named in a Root Path, with no row of its own
    children: list["Record"] = field(default_factory=list)

    @property
    def heading(self) -> str:
        return f"{self.code} {self.name}" if self.code else self.name


def _activities(text: str) -> list[str]:
    """"A (task): «position-offset:2» | Start (evStart) | ..." -> the labels,
    start events first, then in position order, end events last."""
    items = []
    for i, part in enumerate(p.strip() for p in text.split(" | ")):
        if not part:
            continue
        m = _OFFSET.search(part)
        label = _OFFSET.sub("", part).strip()
        rank = 0 if "(evStart)" in label else 2 if "(evEnd)" in label else 1
        items.append((rank, int(m.group(1)) if m else 10**6, i, label))
    return [label for *_, label in sorted(items)]


def read(path: Path = SOURCE) -> list[Record]:
    """Every row that names a process or an object, in sheet order. An object
    row with nothing else in it still names a BPMN code the graph links to."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [_fix(c) for c in next(rows)]
        missing = [c for c in COLUMNS.values() if c not in header]
        if missing:
            raise ValueError(f"{path.name} is not a BPML process export: no {', '.join(missing)} column")
        at = {key: header.index(col) for key, col in COLUMNS.items()}
        records = []
        for row in rows:
            cell = {key: _fix(row[i]) if i < len(row) else "" for key, i in at.items()}
            if cell["name"] in _PLACEHOLDERS or cell["name"].startswith("DELETE "):
                continue
            # A zero-width space in a name came out of the export as "?":
            # "Design assets?". No process name asks a question.
            cell["name"] = re.sub(r"(?<=\w)\?(?=\s|$)", "", cell["name"])
            m = CODE_LINE.match(cell["name"])
            path_parts = [s.strip() for s in cell["path"].split(">") if s.strip()]
            while path_parts and path_parts[0] in _PATH_PREFIX:
                path_parts.pop(0)
            record = Record(
                code=m.group(1) if m else "", name=(m.group(2) if m else cell["name"]).strip(),
                path=path_parts, description=cell["description"],
                process_type=cell["process_type"], status=cell["status"],
                accountable=cell["accountable"], roles=cell["roles"],
                inputs=cell["inputs"], outputs=cell["outputs"],
                activities=_activities(cell["activities"]),
            )
            records.append(record)
        return records
    finally:
        wb.close()


def build(records: list[Record]) -> list[Record]:
    """Nest the records into the hierarchy and return its roots.

    A code the export lists twice keeps its first row ("2.0 Acquire to Dispose",
    not the descoped "2.0 A2D"; "9.2.3 Perform Fixed Asset Accounting", not its
    "(Copy)"), with blanks filled from the later one and its activities added:
    the copy of 9.2.3 lists steps the original does not. A code that has no row
    but is named in a Root Path ("8.6.3 Plan and manage fleet movements") gets
    a section of its own, or its sub-processes would hang off nothing. A BPMN
    object row sits under the last numbered process of its Root Path."""
    procs: dict[str, Record] = {}
    objects: list[Record] = []
    for r in records:
        if not r.code:
            objects.append(r)
        elif r.code in procs:
            kept = procs[r.code]
            for attr in ("description", "process_type", "status", "accountable",
                         "roles", "inputs", "outputs"):
                if not getattr(kept, attr) and getattr(r, attr):
                    setattr(kept, attr, getattr(r, attr))
            kept.activities += [a for a in r.activities if a not in kept.activities]
        else:
            procs[r.code] = r

    for r in list(procs.values()) + objects:
        for i, segment in enumerate(r.path):
            m = CODE_LINE.match(segment)
            if m and m.group(1) not in procs:
                procs[m.group(1)] = Record(code=m.group(1), name=m.group(2).strip(),
                                           path=r.path[:i], synthesised=True)

    def path_owner(r: Record) -> str | None:
        return next((m.group(1) for s in reversed(r.path) if (m := CODE_LINE.match(s))), None)

    roots: list[Record] = []
    for code in sorted(procs, key=sort_key):
        parent = parent_of(code)
        if parent not in procs:
            # Some codes do not extend their parent's ("2.4.02.06.01" under
            # "2.4.2.06"). The document nests them by Root Path so their chunks
            # keep the context; bpml.py still derives parents from the codes.
            parent = path_owner(procs[code])
        (procs[parent].children if parent in procs else roots).append(procs[code])
    for r in objects:
        owner = path_owner(r)
        (procs[owner].children if owner in procs else roots).append(r)
    for p in procs.values():
        # Objects detail the process itself, so they come before its sub-processes.
        p.children.sort(key=lambda c: (bool(c.code), sort_key(c.code) if c.code else ()))
    return roots


def _line(text: str) -> str:
    return " ".join(text.split())


def _prose(text: str) -> str:
    """Description text, safe to put in the document: a line that would read
    as a heading, a table or a fence is escaped, and runs of blank lines close."""
    lines = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith(("#", "|", "```", "- **")):
            stripped = "\\" + stripped
        lines.append(stripped)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def render(roots: list[Record]) -> str:
    out = [
        f"# {TITLE}", "",
        "The Solvay process house (BPML), exported from Signavio: one section per "
        "process, nested under its parent. Each section gives the process's level, "
        "its path in the house, its type and status, who is responsible, and then "
        "its description and activities.", "",
    ]

    def emit(r: Record, depth: int) -> None:
        out.extend([f"{'#' * min(depth + 1, 6)} {_line(r.heading)}", ""])
        meta = []
        if r.code:
            meta.append(f"- **Level:** {level_of(r.code)}")
        if r.path:
            meta.append(f"- **Path:** {_line(' > '.join(r.path))}")
        meta += [f"- **{label}:** {_line(getattr(r, attr))}" for label, attr in FIELDS if getattr(r, attr)]
        if r.synthesised:
            meta.append("- **Note:** The export has no row for this process; it is named "
                        "in the path of its sub-processes.")
        if meta:
            out.extend(meta + [""])
        if r.description:
            out.extend([_prose(r.description), ""])
        if r.activities:
            out.extend(["**Activities:**", ""])
            # One block each, so the chunker can cut a long list between items.
            for a in r.activities:
                out.extend([f"- {_line(a)}", ""])
        for child in r.children:
            emit(child, depth + 1)

    for root in roots:
        emit(root, 1)
    return "\n".join(out).rstrip() + "\n"


_HEADING = re.compile(r"^#{2,6}\s+(.*?)\s*$")
_FIELD = re.compile(r"^- \*\*(.+?):\*\*\s?(.*)$")
# A lettered BPMN code: "O-020-010-010 Review ...", "PE-040-030 Plan ...".
BPMN_CODE = re.compile(r"^([A-Za-z][A-Za-z0-9]{0,3}-\d{2,3}(?:-\d{2,3})*)\s+(.*)$", re.S)
# An activity's BPMN kind, and any note after it: "Create Class (task): To be
# managed", "Permit process (task) Future: link ...". Only known kinds, so a
# name like "... (R2R)+ SAP LUM" keeps its brackets.
_KIND = re.compile(r"\s*\((task|subProcess|ev[A-Z]\w*|\w*[Gg]ateway)\)(?:\W.*)?$", re.S)


def sections(text: str) -> list[dict]:
    """Every section of a document render() wrote, in document order:
    {"heading", "fields": {label: value}, "description", "activities"}.

    Works on the chunks' text joined back together as well as on the file:
    each chunk keeps its headings, and a paragraph cut between two chunks only
    gains a paragraph break."""
    out: list[dict] = []
    current: dict | None = None
    mode = "fields"
    description: list[str] = []

    def close() -> None:
        if current is not None:
            current["description"] = re.sub(r"\n{3,}", "\n\n", "\n".join(description)).strip()

    for line in text.splitlines():
        stripped = line.strip()
        m = _HEADING.match(stripped)
        if m:
            close()
            current = {"heading": m.group(1), "fields": {}, "activities": []}
            out.append(current)
            description = []
            mode = "fields"
            continue
        if current is None:
            continue
        if stripped == "**Activities:**":
            mode = "activities"
            continue
        if mode == "activities":
            if stripped.startswith("- "):
                current["activities"].append(stripped[2:])
            continue
        if mode == "fields":
            f = _FIELD.match(stripped)
            if f:
                current["fields"][f.group(1)] = f.group(2).strip()
                continue
            if not stripped:
                continue
            mode = "description"
        description.append(stripped.removeprefix("\\"))
    close()
    return out


def parse(text: str) -> list[dict]:
    """The numbered processes, in document order: {"code", "name",
    "description", "process_type", "status", ...}. A BPMN object's section
    belongs to no process of its own and is skipped."""
    out: list[dict] = []
    for section in sections(text):
        m = CODE_LINE.match(section["heading"])
        if not m:
            continue
        record = {"code": m.group(1), "name": m.group(2).strip(), "description": section["description"]}
        for label, attr in FIELDS:
            if label in section["fields"]:
                record[attr] = section["fields"][label]
        out.append(record)
    return out


def hierarchy(text: str) -> dict:
    """The parent and name of every code the document names, numbered and
    lettered: {"parent": {code: parent_code}, "name": {code: name}}. The
    knowledge graph chains a cited process up to its value chain with this.

    A numbered process's parent is the last numbered segment of its Path --
    which also places "2.4.02.06.01" under "2.4.2.06", where its code alone
    would not. A lettered code can appear in several processes, so where it is
    performed wins: an activity it is the task (or sub-process) of first, then
    the process an object section of that code sits under, then a process that
    only refers to it as a start, end or intermediate event. Ties go to the
    first in the document. A code's name comes from the same place as its
    parent: the export labels one code differently in different processes."""
    from collections import defaultdict

    # code -> [(rank, order, parent or None, name)]; the lowest rank wins.
    found: dict[str, list[tuple[int, int, str | None, str]]] = defaultdict(list)
    order = 0
    for section in sections(text):
        order += 1
        path = [p.strip() for p in section["fields"].get("Path", "").split(" > ") if p.strip()]
        numbered = []
        for segment in path:
            m = CODE_LINE.match(segment)
            if m:
                numbered.append(m.group(1))
                # Named in a path only: a name, and the weakest one.
                found[m.group(1)].append((9, order, None, m.group(2).strip()))
        own = CODE_LINE.match(section["heading"])
        if own:
            host = own.group(1)
            found[host].append((0, order, numbered[-1] if numbered else None, own.group(2).strip()))
        else:
            host = numbered[-1] if numbered else None
            obj = BPMN_CODE.match(section["heading"])
            if obj:
                found[obj.group(1)].append((1, order, host, _KIND.sub("", obj.group(2)).strip()))
        for activity in section["activities"]:
            order += 1
            m = BPMN_CODE.match(activity)
            if not m:
                continue
            kind = _KIND.search(m.group(2))
            performed = kind is not None and kind.group(1) in ("task", "subProcess")
            found[m.group(1)].append((0 if performed else 2, order, host, _KIND.sub("", m.group(2)).strip()))

    parent: dict[str, str] = {}
    name: dict[str, str] = {}
    for code, entries in found.items():
        name[code] = min(entries)[3]
        placed = [e for e in entries if e[2] and e[2] != code]
        if placed:
            parent[code] = min(placed)[2]
    return {"parent": parent, "name": name}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--target", type=Path, default=TARGET)
    parser.add_argument("--index", action="store_true", help="re-index the written file")
    args = parser.parse_args()

    roots = build(read(args.source))
    text = render(roots)
    args.target.write_text(text, encoding="utf-8")
    processes = parse(text)
    print(f"Wrote {args.target} ({len(text):,} characters, {len(processes)} numbered processes).")
    if args.index:
        from backend.rag import rag

        result = rag.index_path(args.target, force=True)
        print(f"Indexed: {result}")


if __name__ == "__main__":
    main()
